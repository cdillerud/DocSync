"""
Tests for services/new_item_requests.py — reading the Inside Sales "New Item
Template" (sent to items@gamerpackaging.com) and handing its fields to the
Commercial AI catalog intake.

Pins down:
  1. the form is recognized by its layout (Instructions tab naming items@,
     "New Item Request Form" tabs), one record per filled tab, unused tabs
     ignored, attributes kept by label, the concatenated description read;
  2. form problems are named (description over 100 characters, blank
     classification, an item on two tabs, #VALUE!);
  3. a workbook that isn't a request form is skipped quietly;
  4. off by default; dry run records but never sends; a form is sent once;
  5. the URL and key default from the catalog forward settings; the key is
     never stored.
"""
import io

import mongomock_motor
import openpyxl
import pytest

from services import new_item_requests as nir

KEY = "secret-intake-key-123"


def _tab(wb, title, item_no="", category="", classification="", vendor="", vendor_item="", customer="",
         sales_uom="M", ea_case="", ea_pallet="", pallets_truck="", pallets_ctr="", packout="", dims="",
         attributes=(), description="", notes=""):
    ws = wb.create_sheet(title)
    rows = [
        ("New Item Request Form", None, None, None),
        ("Header ", None, "Character Count", "Notes & Examples"),
        ("Item No.", item_no or None, len(item_no), "Use the Supplier Item # unless the Supplier doesn't have one. "),
        ("Item Category", category or None),
        ("Item Classification", classification or None, None, "Custom – ... Stock – Supplier or Gamer Mold"),
        ("BC Vendor No.", vendor or None),
        ("Vendor Item No.", vendor_item or None),
        ("BC Customer No.", customer or None, None, "If stock, use BC Customer No. not GAMER. "),
        ("Customer Item No.", item_no or None),
        ("Unit Of Measure ", None),
        ("Sales UOM", sales_uom),
        ("Purchase UOM", sales_uom),
        ("Warehouse UOM", "PLT"),
        ("Packout ", None),
        ("EA/Case (only if case packed)", ea_case or None),
        ("EA/Pallet", ea_pallet or None),
        ("Pallets/Truck", pallets_truck or None),
        ("Pallets/Container", pallets_ctr or None),
        ("Packout Extended Text", packout or None, len(packout), "Case Packed - ..."),
        ("Dims & Wt Extended Text", dims or None, len(dims)),
        ("Attributes", None),
    ]
    rows += [(label, value or None, None, "examples") for label, value in attributes]
    rows += [
        ("Additional Notes ", None),
        ("Additional Notes ", notes or None, len(notes)),
        ("Final Extended Text View", None),
        ("Description", description or ", ", len(description)),
        ("Packout Extended Text", packout or 0, 1),
        ("Dims & Wt Extended Text", dims or 0, 1),
        ("Additional Notes ", notes or 0, 1),
    ]
    for r in rows:
        ws.append(list(r))
    return ws


def _workbook(filled=True, extra=None) -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = "Instructions"
    wb["Instructions"].append(["Instructions"])
    wb["Instructions"].append([5, "Save and send to Items@gamerpackaging.com"])
    if filled:
        _tab(wb, "Vessel", item_no="PB16OZBIGFOOT", category="Bottle", classification="Stock",
             vendor="PB16OZBIGFOOT", vendor_item="AMCOR", customer="BEVGROW", ea_pallet="4284/plt", pallets_truck=26,
             packout="4,284/plt, 26 plt/TL, 111,384/TL", dims='56x44x102", 350 lbs/plt',
             attributes=[("Volume", 16), ("Volume UOM", "oz"), ("Finish", "28-1810mm"), ("Finish Type", "Twist"),
                         ("Color", "Clear"), ("Material ", "PET"), ("Shape/Style", "Beverage"),
                         ("Product Type ", "Bottle"), ("Decorartion", None)],
             description="16oz, 28-1810mm, Twist, Clear, PET, Beverage, Bottle", notes="Carbonation apprvd")
        _tab(wb, "20 ft gold closure", item_no="63-2030GoldLug20FT", category="Cap", classification="Stock",
             vendor="THCENT", vendor_item="63-2030GoldLug20FT", customer="GAMER", ea_case=1300,
             pallets_truck="35 cs/plt", pallets_ctr="CTR", packout="1300/cs, 35 cs/plt, 455,000/CTR, 10 plts/20FT CTR",
             attributes=[("Finish", "63-2030mm"), ("Finish Type", "Lug"), ("Material ", "Metal"), ("Color", "Gold")],
             description="63-2030mm, Lug, Metal, Button, Gold, Cap, Plastisol compound, 20x12x16 case dims")
    _tab(wb, "Carriers")  # an unused tab
    if extra:
        extra(wb)
    wb.create_sheet("Dropdown Options").append(["Sale/Purch UOM", "Whse UOM"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_form_read_one_record_per_filled_tab():
    is_form, records, issues = nir.parse_workbook(_workbook())
    assert is_form and issues == []
    assert [r["sheet"] for r in records] == ["Vessel", "20 ft gold closure"]
    vessel = records[0]
    assert vessel["item_no"] == "PB16OZBIGFOOT" and vessel["item_classification"] == "Stock"
    assert vessel["bc_vendor_no"] == "PB16OZBIGFOOT" and vessel["vendor_item_no"] == "AMCOR"   # as typed
    assert vessel["pallets_per_truck"] == "26" and vessel["ea_per_pallet"] == "4284/plt"
    assert vessel["packout_text"] == "4,284/plt, 26 plt/TL, 111,384/TL"
    assert vessel["attributes"]["Volume"] == "16" and vessel["attributes"]["Finish"] == "28-1810mm"
    assert vessel["attributes"]["Material"] == "PET" and "Decorartion" not in vessel["attributes"]
    assert vessel["description"] == "16oz, 28-1810mm, Twist, Clear, PET, Beverage, Bottle"
    assert vessel["additional_notes"] == "Carbonation apprvd" and vessel["issues"] == []
    cap = records[1]
    assert cap["pallets_per_truck"] == "35 cs/plt" and cap["ea_per_case"] == "1300"


def test_form_problems_are_named():
    def more(wb):
        _tab(wb, "Vessel (2)", item_no="PB16OZBIGFOOT", category="Bottle", description="x" * 120)
        _tab(wb, "Other", item_no="10142", category="Pail", classification="Stock", ea_pallet="#VALUE!")
    is_form, records, issues = nir.parse_workbook(_workbook(extra=more))
    dup = [r for r in records if r["sheet"] == "Vessel (2)"][0]
    assert any("120 characters" in i for i in dup["issues"]) and "Item Classification is blank" in dup["issues"]
    assert any("on two tabs" in i for i in issues)
    other = [r for r in records if r["sheet"] == "Other"][0]
    assert any("#VALUE!" in i for i in other["issues"])


def test_other_workbooks_are_not_forms():
    wb = openpyxl.Workbook()
    wb.active.append(["Item", "Price"])
    buf = io.BytesIO()
    wb.save(buf)
    assert nir.parse_workbook(buf.getvalue())[0] is False
    assert nir.parse_workbook(b"not a workbook")[0] is False
    is_form, records, issues = nir.parse_workbook(_workbook(filled=False))
    assert is_form and records == [] and issues == ["no tab has an Item No. filled in"]


def test_cheap_filter():
    assert nir.might_be_request({"file_name": "New Item Template ANGIE.xlsx"}) == (True, "workbook")
    assert nir.might_be_request({"file_name": "scan.pdf"})[0] is False
    assert nir.might_be_request({"file_name": "a.xlsx", "new_item_request_sent_utc": "x"}) == (False, "already_sent")


def test_settings_default_from_the_catalog_forward(monkeypatch):
    monkeypatch.delenv("NEW_ITEM_REQUESTS_URL", raising=False)
    monkeypatch.delenv("NEW_ITEM_REQUESTS_KEY", raising=False)
    monkeypatch.delenv("NEW_ITEM_REQUESTS_ENABLED", raising=False)
    monkeypatch.setenv("CATALOG_FORWARD_URL", "https://x.test/api/v70/catalog/intake/submit")
    monkeypatch.setenv("CATALOG_FORWARD_KEY", KEY)
    s = nir.get_settings()
    assert s.url == "https://x.test/api/v70/catalog/intake/new-items" and s.key == KEY and s.enabled is False
    assert s.public()["key_configured"] is True and KEY not in str(s.public())


# ── sending ────────────────────────────────────────────────────

@pytest.fixture
def fake_db(monkeypatch):
    db = mongomock_motor.AsyncMongoMockClient()["test_new_item_requests"]
    monkeypatch.setattr(nir, "get_db", lambda: db)
    return db


@pytest.fixture
def enabled(monkeypatch):
    monkeypatch.setenv("NEW_ITEM_REQUESTS_ENABLED", "true")
    monkeypatch.setenv("NEW_ITEM_REQUESTS_URL", "https://x.test/api/v70/catalog/intake/new-items")
    monkeypatch.setenv("NEW_ITEM_REQUESTS_KEY", KEY)
    monkeypatch.delenv("NEW_ITEM_REQUESTS_DRY_RUN", raising=False)


@pytest.fixture
def files(monkeypatch):
    by_id = {"form": _workbook(), "plain": b"PK not really"}

    async def _load(doc):
        return by_id.get(doc["id"])
    monkeypatch.setattr(nir, "_load_bytes", _load)


@pytest.fixture
def posts(monkeypatch):
    sent = []

    async def _post(settings, payload):
        sent.append({"key": settings.key, "payload": payload})
        return {"status": "ACCEPTED", "stored": len(payload["records"]), "skipped": []}
    monkeypatch.setattr(nir, "_post", _post)
    return sent


async def _insert(db):
    await db.hub_documents.insert_many([
        {"id": "form", "file_name": "New Item Template ANGIE.xlsx", "email_subject": "New items",
         "email_sender": "rep@gamerpackaging.com", "created_utc": "2026-10-22T20:00:00+00:00"},
        {"id": "plain", "file_name": "Pretium price list.xlsx", "created_utc": "2026-10-22T20:00:00+00:00"},
        {"id": "pdf", "file_name": "invoice.pdf"},
    ])


@pytest.mark.asyncio
async def test_off_by_default(fake_db, files, posts, monkeypatch):
    monkeypatch.delenv("NEW_ITEM_REQUESTS_ENABLED", raising=False)
    await _insert(fake_db)
    out = await nir.process_documents(["form"])
    assert out["reason"] == "disabled" and posts == []


@pytest.mark.asyncio
async def test_sends_the_form_once_and_skips_the_rest(fake_db, enabled, files, posts):
    await _insert(fake_db)
    out = await nir.process_documents(["form", "plain", "pdf"])
    assert out["counts"] == {"sent": 1, "skipped": 1}
    assert len(posts) == 1 and posts[0]["key"] == KEY
    payload = posts[0]["payload"]
    assert payload["file_name"] == "New Item Template ANGIE.xlsx" and payload["sender"] == "rep@gamerpackaging.com"
    assert [r["item_no"] for r in payload["records"]] == ["PB16OZBIGFOOT", "63-2030GoldLug20FT"]
    doc = await fake_db.hub_documents.find_one({"id": "form"})
    assert doc["new_item_request"]["status"] == "sent" and doc["new_item_request_sent_utc"]
    assert KEY not in str(doc) and KEY not in str(await fake_db[nir.LOG_COLLECTION].find({}).to_list(10))
    again = await nir.process_documents(["form"])
    assert again["counts"] == {"skipped": 1} and len(posts) == 1


@pytest.mark.asyncio
async def test_dry_run_records_but_sends_nothing(fake_db, enabled, files, posts, monkeypatch):
    monkeypatch.setenv("NEW_ITEM_REQUESTS_DRY_RUN", "true")
    await _insert(fake_db)
    out = await nir.process_documents(["form"])
    assert out["counts"] == {"dry_run": 1} and posts == []
    doc = await fake_db.hub_documents.find_one({"id": "form"})
    assert doc["new_item_request"]["records"] == 2 and not doc.get("new_item_request_sent_utc")


@pytest.mark.asyncio
async def test_failed_send_is_recorded_and_does_not_raise(fake_db, enabled, files, monkeypatch):
    async def _boom(settings, payload):
        raise RuntimeError("catalog intake returned HTTP 500")
    monkeypatch.setattr(nir, "_post", _boom)
    await _insert(fake_db)
    out = await nir.process_documents(["form"])
    assert out["counts"] == {"failed": 1}
    doc = await fake_db.hub_documents.find_one({"id": "form"})
    assert doc["new_item_request"]["status"] == "failed" and "HTTP 500" in doc["new_item_request"]["reason"]


@pytest.mark.asyncio
async def test_backfill_dry_run_works_while_off(fake_db, files, posts, monkeypatch):
    monkeypatch.delenv("NEW_ITEM_REQUESTS_ENABLED", raising=False)
    await _insert(fake_db)
    out = await nir.backfill(days=1095, dry_run=True)
    assert out["candidates"] == 1 and out["counts"] == {"dry_run": 1} and posts == []
    assert (await nir.backfill(dry_run=False))["enabled"] is False
